"""
FastAPI bridge between the React frontend (case-compass / Lex Harvester) and
the Python pipeline (store + search + extract).

Routes:
  GET  /api/stats           - document/chunk/extraction counts
  GET  /api/statutes        - search-backed, capped at 50. Reads the
                              curated multi-state Postgres (port 5433).
                              Query params: q, jurisdiction (CSV), limit.
  POST /api/search          - hybrid (semantic + keyword) search,
                              with a section-number short-circuit
  POST /api/chat            - Claude-backed chat reply. Delegates to
                              api.chat_wrapper.ChatWrapper (Opus 4.7 +
                              adaptive thinking + tool-use retrieval +
                              prompt caching).
  POST /api/upload          - extract text from an uploaded file (PDF /
                              DOCX / TXT). Optional `ingest=true` to also
                              add the file to the searchable corpus.

Run:
    uv run uvicorn api.main:app --reload --port 8000
"""
from __future__ import annotations

import json
import os
import re
from contextlib import contextmanager
from pathlib import Path
from typing import Any, Optional

from dotenv import load_dotenv
from fastapi import FastAPI, HTTPException, UploadFile, File, Form
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field

from store.db import Database
from search.semantic import SemanticIndex


load_dotenv()

DB_PATH = os.environ.get("HACKATHON_DB", "hackathon.db")
STATUTES_DSN = (
    os.environ.get("POSTGRES_DSN")
    or os.environ.get("DATABASE_URL")
    or "postgresql://postgres:postgres@localhost:5433/new_api"
)

app = FastAPI(title="Lex Harvester API", version="0.1.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],  # dev-only; tighten before prod
    allow_methods=["*"],
    allow_headers=["*"],
)


_db: Optional[Database] = None
_idx: Optional[SemanticIndex] = None


def get_db() -> Database:
    global _db
    if _db is None:
        _db = Database(DB_PATH)
        _db.init_schema()
    return _db


def get_index() -> SemanticIndex:
    global _idx
    if _idx is None:
        _idx = SemanticIndex(get_db())
    return _idx


# ---------- Schemas (mirror src/lib/statutes.ts on the frontend) ----------


class CaseRef(BaseModel):
    name: str
    citation: str


class Source(BaseModel):
    publisher: str
    url: str


class Statute(BaseModel):
    id: str
    jurisdiction: str
    jurisdictionLabel: str
    code: str
    section: str
    title: str
    summary: str
    text: str
    factors: list[str] = Field(default_factory=list)
    related: list[str] = Field(default_factory=list)
    cases: list[CaseRef] = Field(default_factory=list)
    source: Source
    lastVerified: str


class SearchHit(BaseModel):
    """A semantic-search hit, formatted to slot into the UI's result list."""
    id: str
    doc_id: int
    chunk_idx: int
    score: float
    text: str
    char_start: int
    char_end: int
    source_url: str
    metadata: dict = Field(default_factory=dict)


class SearchRequest(BaseModel):
    query: str
    top_k: int = 10
    # Optional list of jurisdiction codes (e.g. ["CA", "TX"]) to scope the
    # search to. When omitted/empty, all jurisdictions are searched.
    jurisdictions: list[str] = Field(default_factory=list)


class ChatMessage(BaseModel):
    role: str  # "user" | "assistant"
    text: str


class AttachedFile(BaseModel):
    filename: str
    text: str


class ChatRequest(BaseModel):
    message: str
    history: list[ChatMessage] = Field(default_factory=list)
    matter_name: Optional[str] = None
    matter_caption: Optional[str] = None
    attached_files: list[AttachedFile] = Field(default_factory=list)


class ChatResponse(BaseModel):
    text: str
    # Statutes the model actually pulled this turn via search_statutes /
    # get_statute. Empty when no statute tool ran. The FE renders these as
    # cards under the assistant reply.
    statutes: list[Statute] = Field(default_factory=list)


class UploadResponse(BaseModel):
    filename: str
    size: int
    char_count: int
    text: str
    ingested: bool = False
    doc_id: Optional[int] = None
    chunks: int = 0


class StatuteList(BaseModel):
    """Response shape for /api/statutes — list + total match count.

    `items` is capped server-side (currently 50). `total` is the number of
    rows that match the query across the whole table, so the FE can show
    "50 of 1,234".
    """
    items: list[Statute]
    total: int


class Comparable(BaseModel):
    """A damages comparable — a past PI verdict/settlement an attorney can sort against."""
    id: str
    case_name: str
    citation: str
    jurisdiction: str
    jurisdictionLabel: str
    year: Optional[int] = None
    kind: str = "verdict"  # "verdict" | "settlement"
    plaintiff: Optional[str] = None
    defendant: Optional[str] = None
    factors: list[str] = Field(default_factory=list)
    injuries: list[str] = Field(default_factory=list)
    fact_pattern: str = ""
    award_total_usd: float = 0
    economic_usd: float = 0
    non_economic_usd: float = 0
    punitive_usd: float = 0
    source_url: str = ""
    source_quote: str = ""


# ---------- Statute sourcing (real DB, with seed fallback) ----------
# Lives server-side so the UI hits a real API, even before any docs are ingested.

_STATUTE_SEED_PATH = Path(__file__).parent / "statutes_seed.json"
def _load_seed_statutes() -> list[dict]:
    if _STATUTE_SEED_PATH.exists():
        return json.loads(_STATUTE_SEED_PATH.read_text(encoding="utf-8"))
    return []


def _verdict_extraction_to_comparable(ext: dict, doc: dict) -> Optional[dict]:
    """Map a stored Verdict extraction + its parent opinion doc into the
    Comparable shape. Returns None if the extraction lacks a damages amount."""
    data = ext.get("data") or {}
    total = data.get("total_amount_usd") or 0
    if not total:
        return None

    meta = doc.get("metadata") or {}
    state_raw = (meta.get("jurisdiction") or "").upper().strip()
    state_label = (meta.get("jurisdictionLabel") or meta.get("court") or "").strip()
    case_name = (meta.get("case_name") or "").strip() or "Unknown case"
    citation = (meta.get("citation") or "").strip()
    date_filed = (meta.get("date_filed") or "")[:10]
    year = None
    if len(date_filed) >= 4 and date_filed[:4].isdigit():
        year = int(date_filed[:4])
    elif data.get("decision_date"):
        try:
            year = int(str(data["decision_date"])[:4])
        except (TypeError, ValueError):
            pass

    award_type = (data.get("award_type") or "").lower().strip()
    kind = "settlement" if "settle" in award_type else "verdict"

    return {
        "id": f"opinion-{doc.get('id')}-ext-{ext.get('id')}",
        "case_name": case_name,
        "citation": citation or f"({date_filed})" if date_filed else case_name,
        "jurisdiction": state_raw or "??",
        "jurisdictionLabel": state_label or state_raw or "Unknown",
        "year": year,
        "kind": kind,
        "plaintiff": data.get("plaintiff"),
        "defendant": data.get("defendant"),
        "factors": [],
        "injuries": [],
        "fact_pattern": "",
        "award_total_usd": float(total),
        "economic_usd": float(data.get("compensatory_amount_usd") or 0),
        "non_economic_usd": 0,
        "punitive_usd": float(data.get("punitive_amount_usd") or 0),
        "source_url": doc.get("source_url", ""),
        "source_quote": ext.get("source_quote") or "",
    }


def _load_comparables_from_db() -> list[dict]:
    """Pull stored Verdict extractions and shape them into Comparables."""
    db = get_db()
    extractions = db.get_extractions_by_schema("Verdict")
    if not extractions:
        return []

    out: list[dict] = []
    for ext in extractions:
        doc = db.get_document(ext.get("doc_id"))
        if not doc:
            continue
        comp = _verdict_extraction_to_comparable(ext, doc)
        if comp:
            out.append(comp)
    return out


def _load_comparables() -> list[dict]:
    """Comparables come exclusively from real Verdict extractions in the DB
    (populated by scripts/ingest_comparables.py over CourtListener opinions).
    No hand-curated seed.
    """
    return _load_comparables_from_db()


_JURISDICTION_CODES = {
    "California": "CA",
    "New York": "NY",
    "Texas": "TX",
    "Florida": "FL",
    "Illinois": "IL",
    "Washington": "WA",
    "Arizona": "AZ",
    "Georgia": "GA",
}


def _doc_to_statute(doc: dict) -> dict:
    """Shape a documents-table row into the frontend's Statute type."""
    meta = doc.get("metadata") or {}
    state = (meta.get("state") or "").strip()
    section = (meta.get("section") or "").strip()
    citation_full = (meta.get("citation") or "").strip()
    statute_name = (meta.get("statute_name") or "").strip()
    gold = (meta.get("category_gold") or "").strip()
    is_upload = bool(meta.get("upload"))
    upload_filename = (meta.get("filename") or "").strip()

    if is_upload:
        # User-uploaded files don't have legal-citation metadata; surface them
        # in the list keyed by filename so they're discoverable to the chat.
        jurisdiction_code = "DOC"
        code_name = "Uploaded Document"
        title = upload_filename or "Uploaded document"
    else:
        jurisdiction_code = _JURISDICTION_CODES.get(state, (state[:2].upper() or "??"))
        code_name = "Vehicle Code" if "Veh" in citation_full else (citation_full or "Code")
        title = gold or statute_name or f"§ {section}"

    raw_text = doc.get("raw_text") or ""
    summary = raw_text.strip().split("\n")[0][:240]

    factors = [gold] if gold else []

    # Pull predicted-tag extraction(s) so the UI can show model output too.
    extractions = doc.get("_extractions") or []
    for ext in extractions:
        data = ext.get("data") or {}
        cat = data.get("primary_category")
        if cat and cat not in factors:
            factors.append(cat)

    if is_upload:
        stable_id = f"upload-{doc.get('id')}"
        jurisdiction_label = "Uploaded"
        publisher = "User upload"
        section_for_card = upload_filename or "—"
        if not factors:
            factors = ["Uploaded Document"]
    else:
        section_slug = (
            section.replace("(", "-").replace(")", "").replace(".", "-").lower()
            or str(doc.get("id"))
        )
        stable_id = f"{jurisdiction_code.lower()}-vc-{section_slug}"
        jurisdiction_label = state or jurisdiction_code
        publisher = (
            "California Legislative Information" if "Cal" in citation_full else "Legislature"
        )
        section_for_card = section

    return {
        "id": stable_id,
        "jurisdiction": jurisdiction_code,
        "jurisdictionLabel": jurisdiction_label,
        "code": code_name,
        "section": section_for_card,
        "title": title,
        "summary": summary,
        "text": raw_text,
        "factors": factors,
        "related": [],
        "cases": [],
        "source": {
            "publisher": publisher,
            "url": doc.get("source_url", ""),
        },
        "lastVerified": (doc.get("fetched_at") or "")[:10],
    }


def _is_statute_doc(source_url: str, metadata: dict) -> bool:
    """True iff this document represents an actual statute the UI should show
    on its Statute panel. Excludes uploaded files and CourtListener verdict
    opinions, both of which live in the same `documents` table."""
    if (source_url or "").startswith("upload://"):
        return False
    if (metadata or {}).get("kind") == "verdict":
        return False
    if (metadata or {}).get("upload"):
        return False
    return True


def _load_statutes_from_db() -> list[dict]:
    """Pull statute documents from SQLite and shape them into Statute objects.
    Uploads and verdict opinions live in the same table but are excluded —
    they have their own surfaces (upload route + Comparables)."""
    db = get_db()
    with db.conn() as c:
        rows = c.execute(
            "SELECT id, source_url, raw_text, metadata, fetched_at "
            "FROM documents ORDER BY id"
        ).fetchall()

    if not rows:
        return []

    # Hydrate extractions in one query per doc (corpus is small; fine here).
    out = []
    for row in rows:
        d = dict(row)
        d["metadata"] = json.loads(d["metadata"]) if d["metadata"] else {}
        if not _is_statute_doc(d["source_url"], d["metadata"]):
            continue
        d["_extractions"] = db.get_extractions_for_doc(d["id"])
        out.append(_doc_to_statute(d))
    return out


# ---------- Postgres-backed statute search (curated multi-state corpus) ----------


@contextmanager
def _statutes_pg_conn():
    """Open a short-lived connection to the curated statutes Postgres (5433)."""
    import psycopg
    from psycopg.rows import dict_row

    conn = psycopg.connect(STATUTES_DSN, row_factory=dict_row)
    try:
        yield conn
    finally:
        conn.close()


_PG_PUBLISHER_BY_JURISDICTION = {
    "CA": "California Legislative Information",
    "NY": "New York State Senate",
    "TX": "Texas Legislature Online",
    "CO": "Colorado General Assembly",
    "FL": "The Florida Senate",
    "NV": "Nevada Legislature",
    "OR": "Oregon State Legislature",
}


def _pg_row_to_statute(row: dict) -> dict:
    """Map a row from the Postgres `statutes` table into the FE Statute shape."""
    citation = (row.get("canonical_citation") or "").strip()
    title = (row.get("title") or "").strip() or citation or "(untitled)"
    full = (row.get("complete_statute") or "").strip()

    summary = (row.get("plain_english_summary") or "").strip()
    if not summary:
        summary = full[:240]
        if len(full) > 240:
            summary += "…"

    jurisdiction_code = (row.get("jurisdiction_code") or "").strip()
    publisher = (
        (row.get("source_name") or "").strip()
        or _PG_PUBLISHER_BY_JURISDICTION.get(jurisdiction_code)
        or "Legislature"
    )

    updated_at = row.get("updated_at")
    last_verified = str(updated_at).split(" ")[0] if updated_at else ""

    return {
        "id": str(row.get("id")),
        "jurisdiction": jurisdiction_code,
        "jurisdictionLabel": (row.get("jurisdiction_name") or jurisdiction_code).strip(),
        "code": (row.get("code_name") or "").strip(),
        "section": (row.get("section_number") or "").strip(),
        "title": title,
        "summary": summary,
        "text": full,
        # Populated by new_api/scripts/tag_statute_factors.py. Empty list when
        # the row hasn't been tagged yet.
        "factors": list(row.get("factors") or []),
        "related": [],
        "cases": [],
        "source": {
            "publisher": publisher,
            "url": (row.get("source_url") or "").strip(),
        },
        "lastVerified": last_verified,
    }


_PG_HARD_LIMIT = 50


def _search_statutes_pg(
    q: Optional[str],
    jurisdiction_csv: Optional[str],
    factors_csv: Optional[str],
    limit: int,
) -> tuple[list[dict], int]:
    """Multi-keyword ILIKE search, capped at 50 rows. Empty `q` → top-N default.

    Filters (all optional, AND-combined with each other; any-of within each):
      - q              free-text tokens AND-ed across title / citation / text
      - jurisdiction   CSV of 2-letter codes (CA,NY,...)
      - factors        CSV of contributing-factor strings; matches via
                       Postgres array overlap (``factors && %s::text[]``)

    Returns (items, total) where `total` is the count of statutes matching
    the WHERE clause across the full table — i.e. the number you'd get
    without the LIMIT. Computed via `COUNT(*) OVER ()` so it costs one extra
    aggregate over the same rows scanned, no second round-trip.
    """
    limit = max(1, min(int(limit or _PG_HARD_LIMIT), _PG_HARD_LIMIT))

    where_parts: list[str] = []
    params: list[Any] = []

    tokens = [t for t in re.split(r"\s+", (q or "").strip()) if t]
    for tok in tokens:
        like = f"%{tok}%"
        where_parts.append(
            "(title ILIKE %s OR canonical_citation ILIKE %s OR complete_statute ILIKE %s)"
        )
        params.extend([like, like, like])

    juris_codes = [
        j.strip().upper()
        for j in (jurisdiction_csv or "").split(",")
        if j.strip()
    ]
    if juris_codes:
        placeholders = ",".join(["%s"] * len(juris_codes))
        where_parts.append(f"jurisdiction_code IN ({placeholders})")
        params.extend(juris_codes)

    factor_list = [
        f.strip()
        for f in (factors_csv or "").split(",")
        if f.strip()
    ]
    if factor_list:
        where_parts.append("factors && %s::text[]")
        params.append(factor_list)

    where_sql = (" AND ".join(where_parts)) if where_parts else "TRUE"
    params.append(limit)

    sql = f"""
        SELECT
            id,
            jurisdiction_code,
            jurisdiction_name,
            code_name,
            section_number,
            canonical_citation,
            title,
            complete_statute,
            plain_english_summary,
            source_url,
            source_name,
            factors,
            updated_at,
            COUNT(*) OVER () AS total_count
        FROM statutes
        WHERE {where_sql}
        ORDER BY jurisdiction_code, section_number
        LIMIT %s
    """

    with _statutes_pg_conn() as c:
        rows = c.execute(sql, params).fetchall()

    if not rows:
        # No matches → run a quick COUNT-only query so the FE can confirm "0 of 0"
        # without showing stale total from a prior search. Cheap on this index.
        count_sql = f"SELECT COUNT(*) AS n FROM statutes WHERE {where_sql}"
        with _statutes_pg_conn() as c:
            total = int(c.execute(count_sql, params[:-1]).fetchone()["n"])
        return [], total

    total = int(rows[0]["total_count"])
    items = [_pg_row_to_statute(dict(r)) for r in rows]
    return items, total


# ---------- Routes ----------


@app.get("/api/health")
def health():
    return {"ok": True}


@app.get("/api/stats")
def stats():
    return get_db().stats()


@app.get("/api/statutes", response_model=StatuteList)
def list_statutes(
    q: Optional[str] = None,
    jurisdiction: Optional[str] = None,
    factors: Optional[str] = None,
    limit: int = 50,
):
    """Search the curated multi-state statute Postgres. Capped at 50 rows.

    Query params:
      q             - whitespace-separated keywords (AND-ed). Empty = top-N default.
      jurisdiction  - CSV of 2-letter codes (CA,NY,TX,CO,FL,NV,OR).
      factors       - CSV of contributing-factor strings (any-of via array overlap).
      limit         - max rows, hard-capped at 50.

    Returns ``{items: Statute[], total: int}`` so the FE can render
    "50 of 1,234". Falls back to the SQLite list (then seed JSON) only if
    Postgres is unreachable, so the UI keeps working when the curated DB
    is down.
    """
    try:
        items, total = _search_statutes_pg(q, jurisdiction, factors, limit)
        return {"items": items, "total": total}
    except Exception:
        # Don't take the UI down when the curated DB isn't running.
        real = _load_statutes_from_db()
        fallback = real if real else _load_seed_statutes()
        return {"items": fallback, "total": len(fallback)}


@app.get("/api/comparables", response_model=list[Comparable])
def list_comparables():
    """Damages comparables — past PI verdicts/settlements an attorney can sort against.

    Filtering and sorting are intentionally client-side: the corpus is small,
    the UI wants snappy interactions, and TanStack Query caches the whole list
    on first load.
    """
    return _load_comparables()


_SECTION_QUERY_RE = re.compile(r"^\s*\d+(\.\d+)?(\([a-z0-9]\))?\s*$", re.I)


def _is_section_query(q: str) -> bool:
    return bool(_SECTION_QUERY_RE.match(q or ""))


def _section_lookup(db: Database, query: str) -> list[dict]:
    """Direct lookup of documents whose metadata.section matches a numeric query."""
    base = query.strip().split("(")[0].strip()
    with db.conn() as c:
        rows = c.execute("SELECT id, metadata FROM documents").fetchall()
    hits = []
    seen = set()
    for row in rows:
        meta = json.loads(row["metadata"]) if row["metadata"] else {}
        section = (meta.get("section") or "").split("(")[0].strip()
        if section != base or row["id"] in seen:
            continue
        seen.add(row["id"])
        chunks = db.get_chunks_for_doc(row["id"])
        if not chunks:
            continue
        c0 = chunks[0]
        hits.append({
            "id": c0["id"],
            "doc_id": row["id"],
            "text": c0["text"],
            "score": 1.0,
            "char_start": c0["char_start"],
            "char_end": c0["char_end"],
            "chunk_idx": c0["chunk_idx"],
        })
    return hits


def _stem(w: str) -> str:
    """Cheap suffix stripping so 'driving' matches 'drive'."""
    w = w.lower()
    for suf in ("ing", "ed", "es", "s"):
        if w.endswith(suf) and len(w) > len(suf) + 2:
            return w[: -len(suf)]
    return w


_STOPWORDS = {"the", "a", "an", "of", "and", "or", "to", "for", "in", "on", "at", "is", "be"}


def _keyword_score(query: str, text: str) -> float:
    """Fraction of meaningful query words whose stem appears in text. Range [0, 1]."""
    words = [w.lower() for w in re.findall(r"\w{3,}", query)]
    words = [w for w in words if w not in _STOPWORDS]
    if not words:
        return 0.0
    t = text.lower()
    matches = sum(1 for w in words if _stem(w) in t)
    return matches / len(words)


@app.post("/api/search", response_model=list[SearchHit])
def search(req: SearchRequest):
    if not req.query.strip():
        return []

    db = get_db()
    jur_filter = {j.upper() for j in (req.jurisdictions or []) if j}

    def _passes_jurisdiction(meta: dict) -> bool:
        if not jur_filter:
            return True
        return (meta or {}).get("jurisdiction", "").upper() in jur_filter

    # Section-number queries (e.g. "22107", "21451(a)") — short-circuit to a
    # direct lookup. Embeddings have no idea what those digits mean.
    if _is_section_query(req.query):
        direct = _section_lookup(db, req.query)
        if direct:
            out: list[dict] = []
            for h in direct:
                doc = db.get_document(h["doc_id"]) or {}
                meta = doc.get("metadata") or {}
                if not _passes_jurisdiction(meta):
                    continue
                out.append({
                    "id": f"chunk-{h['id']}",
                    "doc_id": h["doc_id"],
                    "chunk_idx": h["chunk_idx"],
                    "score": h["score"],
                    "text": h["text"],
                    "char_start": h["char_start"],
                    "char_end": h["char_end"],
                    "source_url": doc.get("source_url", ""),
                    "metadata": meta,
                })
            if out:
                return out
        # If section not in corpus, fall through to semantic so the user still
        # gets *something*.

    # Hybrid: pull more candidates than requested, then rerank with a keyword
    # boost. With ~37 short docs, MiniLM scores cluster tightly and a literal-
    # keyword signal is what separates "DUI" from "Improper Passing" for the
    # query "driving under the influence".
    # Widen the candidate pool when a jurisdiction filter is active so the
    # filter has enough qualifying chunks to choose from after dropping
    # off-state matches.
    candidate_mult = 30 if jur_filter else 3
    idx = get_index()
    semantic = idx.search(req.query, top_k=max(req.top_k * candidate_mult, 30))

    KW_WEIGHT = 0.6
    for h in semantic:
        h["_semantic"] = h["score"]
        kw = _keyword_score(req.query, h["text"])
        h["score"] = h["score"] + KW_WEIGHT * kw
    semantic.sort(key=lambda h: -h["score"])

    # Filter to statute-only chunks (no uploads/verdicts) AND to the requested
    # jurisdictions, before truncating to top_k.
    out: list[dict] = []
    for h in semantic:
        doc = db.get_document(h["doc_id"]) or {}
        meta = doc.get("metadata") or {}
        if not _is_statute_doc(doc.get("source_url", ""), meta):
            continue
        if not _passes_jurisdiction(meta):
            continue
        out.append({
            "id": f"chunk-{h['id']}",
            "doc_id": h["doc_id"],
            "chunk_idx": h["chunk_idx"],
            "score": h["score"],
            "text": h["text"],
            "char_start": h["char_start"],
            "char_end": h["char_end"],
            "source_url": doc.get("source_url", ""),
            "metadata": meta,
        })
        if len(out) >= req.top_k:
            break
    return out


_chat_wrapper: Optional["ChatWrapper"] = None


def get_chat_wrapper() -> "ChatWrapper":
    """Lazy singleton — building this loads the embedding model, so do it once."""
    global _chat_wrapper
    if _chat_wrapper is None:
        from api.chat_wrapper import ChatWrapper

        _chat_wrapper = ChatWrapper(index=get_index(), db=get_db())
    return _chat_wrapper


@app.post("/api/chat", response_model=ChatResponse)
def chat(req: ChatRequest):
    api_key = os.environ.get("ANTHROPIC_API_KEY")
    if not api_key or api_key.startswith("replace"):
        raise HTTPException(
            status_code=503,
            detail="ANTHROPIC_API_KEY not set; chat is disabled.",
        )

    wrapper = get_chat_wrapper()
    result = wrapper.reply(
        message=req.message,
        history=[{"role": m.role, "text": m.text} for m in req.history],
        attached_files=[
            {"filename": af.filename, "text": af.text} for af in req.attached_files
        ],
        matter_name=req.matter_name,
        matter_caption=req.matter_caption,
    )
    return ChatResponse(text=result["text"], statutes=result.get("statutes") or [])


# ---------- File upload ----------

_PDF_EXTS = {".pdf"}
_DOCX_EXTS = {".docx"}
_TXT_EXTS = {".txt", ".md", ".markdown", ".rst", ".log", ".csv"}
_MAX_UPLOAD_BYTES = 25 * 1024 * 1024  # 25 MB


def _extract_text(filename: str, content: bytes) -> str:
    """Return plain text from an uploaded file. Raises HTTPException on failure."""
    name = (filename or "").lower()
    suffix = Path(name).suffix

    if suffix in _PDF_EXTS:
        import io
        from pypdf import PdfReader
        try:
            reader = PdfReader(io.BytesIO(content))
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Could not read PDF: {e}")
        pages: list[str] = []
        for i, page in enumerate(reader.pages, start=1):
            try:
                txt = page.extract_text() or ""
            except Exception:
                txt = ""
            if txt.strip():
                pages.append(f"--- Page {i} ---\n{txt.strip()}")
        return "\n\n".join(pages).strip()

    if suffix in _DOCX_EXTS:
        import io
        from docx import Document
        try:
            doc = Document(io.BytesIO(content))
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Could not read DOCX: {e}")
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n".join(paragraphs).strip()

    if suffix in _TXT_EXTS or not suffix:
        for enc in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
            try:
                return content.decode(enc).strip()
            except UnicodeDecodeError:
                continue
        raise HTTPException(status_code=400, detail="Could not decode text file.")

    raise HTTPException(
        status_code=415,
        detail=f"Unsupported file type: {suffix or '(no extension)'}. "
               "Supported: .pdf, .docx, .txt, .md.",
    )


@app.post("/api/upload", response_model=UploadResponse)
async def upload(file: UploadFile = File(...), ingest: bool = Form(False)):
    raw = await file.read()
    if len(raw) > _MAX_UPLOAD_BYTES:
        raise HTTPException(
            status_code=413,
            detail=f"File exceeds {_MAX_UPLOAD_BYTES // (1024 * 1024)} MB limit.",
        )
    if not raw:
        raise HTTPException(status_code=400, detail="Empty file.")

    text = _extract_text(file.filename or "upload", raw)
    if not text:
        raise HTTPException(
            status_code=422,
            detail="No extractable text found (scanned PDF or empty document).",
        )

    resp = UploadResponse(
        filename=file.filename or "upload",
        size=len(raw),
        char_count=len(text),
        text=text,
    )

    if ingest:
        db = get_db()
        idx = get_index()
        source_url = f"upload://{file.filename or 'upload'}"
        doc_id = db.insert_document(
            source_url=source_url,
            raw_text=text,
            metadata={"upload": True, "filename": file.filename, "size": len(raw)},
        )
        # Re-index: remove existing chunks for this doc first.
        with db.conn() as c:
            c.execute("DELETE FROM chunks WHERE doc_id = ?", (doc_id,))
        n_chunks = idx.index_document(doc_id, text)
        resp.ingested = True
        resp.doc_id = doc_id
        resp.chunks = n_chunks

    return resp
